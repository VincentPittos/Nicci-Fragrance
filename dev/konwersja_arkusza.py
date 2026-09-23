#!/usr/bin/env python3
"""
Konwersja bazy wiedzy doradcy (Nicci_Fragrance_Katalog_Produktow.xlsx) na zakładki
Produkty i Zestawy w schemacie, który czyta apps-script/Code.gs.

Skrypt NIE zgaduje po cichu. Każde pole, które trzeba było zmapować, uprościć albo
zaproponować, trafia do zakładki "Do weryfikacji" z wartością źródłową i powodem.
Dane w imporcie są propozycją do akceptacji właściciela, nie danymi produkcyjnymi.

Użycie:
    python3 dev/konwersja_arkusza.py [ścieżka do xlsx]

Wynik (katalog dev/dane/):
    import-do-arkusza.xlsx   zakładki Produkty, Zestawy, Do weryfikacji (do importu w Arkuszach Google)
    produkty.csv, zestawy.csv, do-weryfikacji.csv   te same dane w formie czytelnej w gicie
    arkusz.json              wartości zakładek jak z getValues(), wejście dla dev/zbuduj_mock.js
    raport-konwersji.md      podsumowanie mapowań

Kolumny wewnętrzne źródła (Notatka doradcy, Pewność danych, Uwagi) nie są przenoszone,
bo repozytorium jest publiczne. Pewność danych wpływa tylko na flagi i pole "aktywny".
"""
import csv
import json
import os
import re
import sys
import unicodedata
from collections import Counter, OrderedDict

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_DEFAULT = os.path.join(ROOT, 'refs', 'prywatne', 'Nicci_Fragrance_Katalog_Produktow.xlsx')
OUT = os.path.join(ROOT, 'dev', 'dane')

KOLUMNY_PRODUKTY = ['id', 'aktywny', 'marka', 'nazwa', 'rodzina', 'profil', 'nuty_glowy', 'nuty_serca',
                    'nuty_bazy', 'opis', 'sezon', 'pora', 'trwalosc', 'projekcja', 'intensywnosc',
                    'cena_5', 'cena_10', 'cena_20', 'ml_dostepne', 'podobne', 'zdjecie_url', 'kolejnosc',
                    'okazja']
KOLUMNY_ZESTAWY = ['id', 'aktywny', 'nazwa', 'opis', 'rodzina', 'cena', 'sklad']
KOLUMNY_WERYFIKACJA = ['id', 'pozycja', 'pole', 'w_zrodle', 'w_imporcie', 'powod']

KANON = ['świeża', 'cytrusowa', 'aromatyczna', 'wodna', 'drzewna', 'skórzana', 'szyprowa',
         'słodka', 'orientalna', 'ambrowa', 'gourmand', 'kwiatowa']


def norm(s):
    s = unicodedata.normalize('NFD', str(s or '').lower())
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    return s.replace('ł', 'l').strip()


# ---------- rodzina ----------
# Człon kanoniczny: pierwszy taki człon w nazwie opisowej wygrywa ("Korzenno-drzewna" -> drzewna).
RDZENIE_KANON = [('cytrus', 'cytrusowa'), ('aromaty', 'aromatyczna'), ('drzew', 'drzewna'),
                 ('skorz', 'skórzana'), ('szypr', 'szyprowa'), ('slod', 'słodka'), ('orient', 'orientalna'),
                 ('oreint', 'orientalna'), ('ambr', 'ambrowa'), ('gourmand', 'gourmand'),
                 ('kwiat', 'kwiatowa'), ('swiez', 'świeża'), ('wodn', 'wodna'), ('morsk', 'wodna')]
# Człony spoza listy: używane tylko, gdy w nazwie nie ma żadnego członu kanonicznego.
RDZENIE_ZASTEPCZE = [('korzen', 'orientalna'), ('agar', 'orientalna'), ('zywic', 'orientalna'),
                     ('kadzid', 'orientalna'), ('dym', 'drzewna'), ('owoc', 'słodka'), ('irys', 'kwiatowa'),
                     ('rozan', 'kwiatowa'), ('ziel', 'świeża'), ('herbac', 'świeża'), ('musuj', 'cytrusowa')]


def mapuj_rodzine(tekst):
    """Zwraca (rodzina_kanoniczna|'', pewne: bool, opis_zrodla)."""
    zrodlo = str(tekst or '').split('/')[0].strip()
    czlony = [norm(c) for c in re.split(r'[-\s]+', zrodlo) if c.strip()]
    for i, c in enumerate(czlony):
        for pref, rodz in RDZENIE_KANON:
            if c.startswith(pref):
                return rodz, i == 0, zrodlo
    for c in czlony:
        for pref, rodz in RDZENIE_ZASTEPCZE:
            if c.startswith(pref):
                return rodz, False, zrodlo
    return '', False, zrodlo


# ---------- profil ----------
# Arkusz nie ma kolumny profilu. Propozycja z ogólnej wiedzy o tym, jak producent
# pozycjonuje linię (seria męska albo niszowa bez podziału). Cała kolumna do weryfikacji.
MESKIE_LP = {1, 4, 11, 12, 14, 16, 19, 33, 51, 52, 53, 54, 55, 56, 57, 58, 59, 60, 61, 62, 63, 64, 65,
             66, 67, 68, 69, 70, 71}


def profil_dla(lp):
    return 'męski' if lp in MESKIE_LP else 'unisex'


# ---------- sezon i pora ----------
PORY_ROKU = [('wiosn', 'wiosna'), ('lat', 'lato'), ('jesien', 'jesień'), ('zim', 'zima')]


def mapuj_sezon(tekst):
    t = norm(tekst)
    if not t:
        return '', False
    if 'caly rok' in t:
        wynik = [p for _, p in PORY_ROKU]
    else:
        wynik = [label for pref, label in PORY_ROKU if re.search(r'\b' + pref, t)]
    uproszczone = bool(re.search(r'najlepiej|wczesn|pozn|upal|chlodn|uniwersal|\(', t))
    return ', '.join(wynik), uproszczone


def mapuj_pore(tekst):
    t = norm(tekst)
    if not t:
        return '', False
    dzien, wieczor = 'dzien' in t, 'wieczor' in t
    if dzien and wieczor:
        return 'uniwersalna', t != 'dzien i wieczor'
    if wieczor:
        return 'wieczór', t not in ('wieczor', 'wieczor, noc')
    if dzien:
        return 'dzień', t != 'dzien'
    return '', True


def intensywnosc_z_projekcji(p):
    if p is None or p == '':
        return ''
    p = float(p)
    return 1 if p <= 2.5 else 2 if p <= 3.5 else 3


# ---------- ceny ----------
CENA_RE = re.compile(r'(\d+)\s*ml\s*[–-]\s*(\d+(?:[.,]\d+)?)\s*zł(\s*\(\?\))?(\s*\(promocja\))?', re.I)


def mapuj_ceny(tekst):
    ceny, flagi = {}, []
    for seg in str(tekst or '').split('|'):
        m = CENA_RE.search(seg)
        if not m:
            continue
        ml, cena, niepewna, promo = int(m.group(1)), float(m.group(2).replace(',', '.')), m.group(3), m.group(4)
        if niepewna:
            flagi.append(f'cena {ml} ml oznaczona w źródle jako niepewna "(?)", wariant pominięty do potwierdzenia')
            continue
        if promo:
            flagi.append('cena oznaczona jako promocja; API nie ma pola na starą cenę, pokażemy ją jako zwykłą')
        ceny[ml] = int(cena) if cena.is_integer() else cena
    return ceny, flagi


# ---------- nazwy -> id ----------
ALIASY = {'renaissance': 7, 'torino 21': 6, 'otlands': 17, 'reflection': 16, 'nouvau monde': 4,
          'saffrano absolu': 49, 'y le parfum': 63, "terre d'hermes": 65, 'one milion': 68,
          'the most wanted': 71}


def indeks_nazw(produkty):
    idx = {}
    for p in produkty:
        if not p['nazwa']:
            continue
        idx[norm(p['nazwa'])] = p['lp']
        bez_nawiasu = re.sub(r'\s*\(.*?\)', '', norm(p['nazwa'])).strip()
        idx.setdefault(bez_nawiasu, p['lp'])
    for k, lp in ALIASY.items():
        idx.setdefault(k, lp)
    return idx


def rozpoznaj_nazwy(tekst, idx):
    """Zwraca (lista lp w kolejności, lista nierozpoznanych fragmentów, użyte aliasy)."""
    znalezione, nierozp, aliasy = [], [], []
    for czesc in re.split(r'[;,]', str(tekst or '')):
        t = ' ' + norm(czesc) + ' '
        if not t.strip():
            continue
        trafione = False
        for klucz in sorted(idx, key=len, reverse=True):
            wzor = r'(?<![\w])' + re.escape(klucz) + r'(?![\w])'
            if re.search(wzor, t):
                lp = idx[klucz]
                if lp not in znalezione:
                    znalezione.append(lp)
                if klucz in ALIASY:
                    aliasy.append(f'{czesc.strip()} -> {klucz}')
                t = re.sub(wzor, ' ', t)
                trafione = True
        if not trafione:
            nierozp.append(czesc.strip())
    return znalezione, nierozp, aliasy


def pid(lp):
    return f'p{int(lp):02d}'


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else SRC_DEFAULT
    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb['Katalog zapachów']
    rows = list(ws.iter_rows(values_only=True))
    H = {h: i for i, h in enumerate(rows[3])}

    def kol(r, nazwa):
        v = r[H[nazwa]]
        return v.strip() if isinstance(v, str) else v

    zrodlo = []
    for r in rows[4:]:
        if not any(r):
            continue
        zrodlo.append({
            'lp': int(kol(r, 'Lp.')), 'marka': kol(r, 'Marka') or '', 'nazwa': kol(r, 'Nazwa') or '',
            'status': kol(r, 'Status') or '', 'rodzina': kol(r, 'Rodzina zapachowa/ składniki przedstawione graficznie'),
            'opis': kol(r, 'Opis produktu (tekst sprzedażowy)') or '', 'glowy': kol(r, 'Nuty głowy') or '',
            'serca': kol(r, 'Nuty serca') or '', 'bazy': kol(r, 'Nuty bazy') or '',
            'sezon': kol(r, 'Pora roku'), 'pora': kol(r, 'Pora dnia'), 'okazja': kol(r, 'Okazja / zastosowanie'),
            'trwalosc': kol(r, 'Trwałość 1-5'), 'projekcja': kol(r, 'Projekcja 1-5'),
            'podobne': kol(r, 'Podobne produkty z naszej oferty'), 'ceny': kol(r, 'Pojemności i ceny'),
            'pewnosc': str(kol(r, 'Pewność danych') or ''), 'uwagi': str(kol(r, 'Uwagi / do weryfikacji') or ''),
        })

    idx = indeks_nazw(zrodlo)
    produkty, weryf = [], []
    statystyka = Counter()
    mapa_rodzin = OrderedDict()

    def flag(rec, pole, w_zrodle, w_imporcie, powod):
        nazwa = f"{rec['marka']} {rec['nazwa']}".strip() or f"wiersz Lp. {rec['lp']} (bez nazwy)"
        weryf.append(OrderedDict(id=pid(rec['lp']), pozycja=nazwa, pole=pole,
                                 w_zrodle='' if w_zrodle is None else str(w_zrodle),
                                 w_imporcie='' if w_imporcie is None else str(w_imporcie), powod=powod))

    for s in zrodlo:
        p = OrderedDict((k, '') for k in KOLUMNY_PRODUKTY)
        p['id'] = pid(s['lp'])
        p['marka'], p['nazwa'] = s['marka'], s['nazwa']
        p['kolejnosc'] = s['lp']
        p['nuty_glowy'], p['nuty_serca'], p['nuty_bazy'] = s['glowy'], s['serca'], s['bazy']
        p['opis'] = s['opis']

        # aktywny
        niedostepny = norm(s['status']).startswith('niedostepny')
        aktywny = 'TAK'
        if not s['nazwa'] or not s['marka']:
            aktywny = 'NIE'
            flag(s, 'marka, nazwa', '', '', 'brak marki i nazwy; pozycja wyłączona do czasu uzupełnienia')
        if s['pewnosc'].startswith('Niska') and '(cena' not in s['pewnosc']:
            aktywny = 'NIE'
            flag(s, 'aktywny', 'pewność danych: niska', 'NIE',
                 'instrukcja arkusza: nie publikować bez weryfikacji piramidy nut')
        if s['lp'] == 13:
            aktywny = 'NIE'
            flag(s, 'nazwa', s['nazwa'], s['nazwa'], 'nazwa do potwierdzenia (w ofercie „Virgin Mountain Water”)')
        elif 'Średnia' in s['pewnosc']:
            flag(s, 'nuty / osiągi', 'pewność danych: średnia', 'bez zmian',
                 'instrukcja arkusza: zweryfikować przed publikacją')
        p['aktywny'] = aktywny
        statystyka['aktywny_' + aktywny] += 1

        # rodzina
        rodz, pewne, opis_zr = mapuj_rodzine(s['rodzina'])
        p['rodzina'] = rodz
        mapa_rodzin.setdefault(opis_zr, [rodz, pewne, 0])
        mapa_rodzin[opis_zr][2] += 1
        if not pewne:
            flag(s, 'rodzina', opis_zr, rodz or '(brak)', 'pierwszy człon spoza 12 rodzin kanonicznych, propozycja')

        # profil
        p['profil'] = profil_dla(s['lp'])

        # sezon, pora
        sezon, upr = mapuj_sezon(s['sezon'])
        p['sezon'] = sezon
        if upr:
            flag(s, 'sezon', s['sezon'], sezon, 'uproszczono do listy pór roku')
        if not sezon:
            flag(s, 'sezon', s['sezon'], '', 'brak w źródle')
        pora, upr = mapuj_pore(s['pora'])
        p['pora'] = pora
        if upr:
            flag(s, 'pora', s['pora'], pora, 'uproszczono do: dzień, wieczór, uniwersalna')
        if not pora:
            flag(s, 'pora', s['pora'], '', 'brak w źródle')

        # trwałość, projekcja, intensywność
        for pole in ('trwalosc', 'projekcja'):
            v = s[pole]
            p[pole] = (int(v) if float(v).is_integer() else v) if isinstance(v, (int, float)) else (v or '')
            if p[pole] == '':
                flag(s, pole, '', '', 'brak w źródle')
        p['intensywnosc'] = intensywnosc_z_projekcji(s['projekcja'])

        # okazja
        if isinstance(s['okazja'], str) and s['okazja'].strip():
            p['okazja'] = ', '.join(x.strip() for x in s['okazja'].split(',') if x.strip())
        else:
            if s['okazja'] not in (None, ''):
                flag(s, 'okazja', s['okazja'], '', 'wartość nie jest listą okazji (błąd w źródle)')

        # ceny
        ceny, flagi = mapuj_ceny(s['ceny'])
        for ml in (5, 10, 20):
            p[f'cena_{ml}'] = ceny.get(ml, '')
        for f in flagi:
            flag(s, 'ceny', s['ceny'], ' | '.join(f'{k} ml {v} zł' for k, v in ceny.items()), f)
        if not ceny and not niedostepny:
            flag(s, 'ceny', s['ceny'], '', 'brak cen')
        statystyka['warianty_' + '/'.join(str(k) for k in sorted(ceny)) if ceny else 'warianty_brak'] += 1

        # stan
        p['ml_dostepne'] = 0 if niedostepny else ''
        if niedostepny:
            flag(s, 'ml_dostepne', s['status'], 0, 'niedostępny w źródle; karta pokaże „Wyprzedane” i podobne')

        # podobne
        lps, nierozp, aliasy = rozpoznaj_nazwy(s['podobne'], idx)
        lps = [x for x in lps if x != s['lp']]
        p['podobne'] = ', '.join(pid(x) for x in lps)
        for n in nierozp:
            flag(s, 'podobne', n, '', 'nazwa nie pasuje do żadnej pozycji katalogu')
        for a in aliasy:
            flag(s, 'podobne', a, '', 'dopasowano po poprawce literówki albo skrótu nazwy')

        # zdjęcie: ujednolicony plik z dev/zdjecia (wersja 800 px; strona dobiera 400 px z tej samej nazwy)
        if os.path.exists(os.path.join(ROOT, 'site', 'img', 'produkty', f"{p['id']}-800.webp")):
            p['zdjecie_url'] = f"/img/produkty/{p['id']}-800.webp"
        elif not niedostepny:
            flag(s, 'zdjecie_url', '', '', 'brak zdjęcia: strona producenta blokuje pobieranie, potrzebne własne albo prasowe')

        produkty.append(p)

    # ---------- zestawy ----------
    wz = wb['Zestawy (Discovery Sets)']
    zrows = list(wz.iter_rows(values_only=True))
    zestawy = []
    ZESTAW_RODZINA = {'SWEET & SPICY SET': 'gourmand'}
    nr = 0
    for r in zrows[3:]:
        if not r or not r[0]:
            continue
        nr += 1
        nazwa, sklad_txt, cena = str(r[0]).strip(), str(r[1] or ''), r[2]
        z = OrderedDict((k, '') for k in KOLUMNY_ZESTAWY)
        z['id'] = f'z{nr:02d}'
        z['nazwa'] = nazwa
        rec = {'lp': 0, 'marka': '', 'nazwa': nazwa}

        def zflag(pole, a, b, powod):
            weryf.append(OrderedDict(id=z['id'], pozycja=nazwa, pole=pole, w_zrodle=str(a), w_imporcie=str(b), powod=powod))

        if isinstance(cena, (int, float)):
            z['cena'] = int(cena) if float(cena).is_integer() else cena
        else:
            m = re.search(r'(\d+(?:[.,]\d+)?)', str(cena or ''))
            z['cena'] = int(float(m.group(1).replace(',', '.'))) if m else ''
            zflag('cena', cena, z['cena'], 'cena zapisana jako tekst')
        tekst_skladu = sklad_txt.split('sugerowane:')[-1]
        lps, nierozp, aliasy = rozpoznaj_nazwy(tekst_skladu.replace('(', ',').replace(')', ','), idx)
        nierozp = [n for n in nierozp if n and not re.fullmatch(r'\s*\d*\s*[x×]?\s*\d*\s*ml\s*', n, re.I)
                   and 'do ustalenia' not in n]
        z['sklad'] = ', '.join(f'{pid(x)}:5' for x in lps)
        z['aktywny'] = 'TAK'
        if 'do ustalenia' in sklad_txt:
            z['aktywny'] = 'NIE'
            zflag('sklad', sklad_txt, z['sklad'], 'skład „do ustalenia”; wpisano propozycję ze źródła, zestaw wyłączony')
        if re.search(r'5x\s*ml', sklad_txt, re.I):
            zflag('sklad', sklad_txt, z['sklad'], 'w źródle „5x ml” bez pojemności; przyjęto 5 ml')
        for n in nierozp:
            zflag('sklad', n, '', 'nazwa nie pasuje do żadnej pozycji katalogu')
        for a in aliasy:
            zflag('sklad', a, '', 'dopasowano po poprawce literówki albo skrótu nazwy')
        if len(lps) != 5:
            zflag('sklad', sklad_txt, z['sklad'], f'rozpoznano {len(lps)} z 5 składników')
        z['rodzina'] = ZESTAW_RODZINA.get(nazwa.upper(), '')
        if z['rodzina']:
            zflag('rodzina', '', z['rodzina'], 'propozycja z opisu zestawu w źródle (gourmand wprost w treści)')
        else:
            zflag('rodzina', '', '', 'brak rodziny: quiz nie zaproponuje tego zestawu')
        zflag('opis', '', '', 'źródło ma tylko opis pozycjonowania dla doradcy; opis dla klienta do napisania')
        zestawy.append(z)

    # ---------- zapis ----------
    os.makedirs(OUT, exist_ok=True)

    def zapisz_csv(nazwa, kolumny, wiersze):
        with open(os.path.join(OUT, nazwa), 'w', newline='', encoding='utf-8') as f:
            w = csv.DictWriter(f, fieldnames=kolumny)
            w.writeheader()
            for x in wiersze:
                w.writerow(x)

    zapisz_csv('produkty.csv', KOLUMNY_PRODUKTY, produkty)
    zapisz_csv('zestawy.csv', KOLUMNY_ZESTAWY, zestawy)
    zapisz_csv('do-weryfikacji.csv', KOLUMNY_WERYFIKACJA, weryf)

    out = openpyxl.Workbook()
    for i, (tytul, kolumny, wiersze) in enumerate([('Produkty', KOLUMNY_PRODUKTY, produkty),
                                                   ('Zestawy', KOLUMNY_ZESTAWY, zestawy),
                                                   ('Do weryfikacji', KOLUMNY_WERYFIKACJA, weryf)]):
        ws2 = out.active if i == 0 else out.create_sheet()
        ws2.title = tytul
        ws2.append(kolumny)
        for x in wiersze:
            ws2.append([x[k] for k in kolumny])
        ws2.freeze_panes = 'A2'
    out.save(os.path.join(OUT, 'import-do-arkusza.xlsx'))

    arkusz = {
        'Produkty': [KOLUMNY_PRODUKTY] + [[x[k] for k in KOLUMNY_PRODUKTY] for x in produkty],
        'Zestawy': [KOLUMNY_ZESTAWY] + [[x[k] for k in KOLUMNY_ZESTAWY] for x in zestawy],
    }
    with open(os.path.join(OUT, 'arkusz.json'), 'w', encoding='utf-8') as f:
        json.dump(arkusz, f, ensure_ascii=False, indent=1)

    # ---------- raport ----------
    rodziny = Counter(p['rodzina'] for p in produkty if p['aktywny'] == 'TAK')
    pola = Counter(w['pole'] for w in weryf)
    L = ['# Raport konwersji arkusza', '',
         'Wygenerowany przez `dev/konwersja_arkusza.py`. Wszystkie wartości oznaczone w zakładce '
         '„Do weryfikacji” są propozycją do akceptacji właściciela.', '',
         f'Produkty: {len(produkty)} (aktywne: {statystyka["aktywny_TAK"]}, wyłączone: {statystyka["aktywny_NIE"]}). '
         f'Zestawy: {len(zestawy)}. Pozycji do weryfikacji: {len(weryf)}.', '',
         '## Rodziny aktywnych produktów', '',
         '| Rodzina | Liczba |', '|---|---|']
    L += [f'| {k or "(brak)"} | {v} |' for k, v in sorted(rodziny.items(), key=lambda x: -x[1])]
    L += ['', '## Mapowanie rodzin opisowych na kanoniczne', '',
          '| W źródle | Propozycja | Pewne | Produktów |', '|---|---|---|---|']
    L += [f'| {k} | {v[0] or "(brak)"} | {"tak" if v[1] else "do weryfikacji"} | {v[2]} |'
          for k, v in sorted(mapa_rodzin.items(), key=lambda x: (x[1][0], x[0]))]
    L += ['', '## Reguły', '',
          '* **rodzina:** pierwszy człon nazwy opisowej, który jest jedną z 12 rodzin kanonicznych. '
          'Gdy żadnego nie ma: korzenna, agarowa, żywiczna, kadzidlana → orientalna; dymna → drzewna; '
          'owocowa → słodka; irysowa, różana → kwiatowa; zielona, herbaciana → świeża.',
          '* **profil:** arkusz go nie ma. Propozycja z ogólnej wiedzy o linii producenta: serie męskie '
          '(np. Sauvage, Eros, Bottled, Aventus) → męski, reszta → unisex. Cała kolumna do weryfikacji.',
          '* **intensywność:** z projekcji 1–5: do 2,5 → 1, do 3,5 → 2, od 4 → 3.',
          '* **sezon:** „cały rok” → wszystkie cztery; kwalifikatory (najlepiej, wczesna, późna) pominięte.',
          '* **pora:** dzień i wieczór razem → uniwersalna.',
          '* **ml_dostepne:** źródło nie ma stanów. Puste = do uzupełnienia przez właściciela, '
          'niedostępne w źródle → 0.',
          '* **aktywny = NIE:** brak nazwy (Lp. 56–58), niska pewność danych, nazwa do potwierdzenia (Lp. 13).',
          '', '## Pola do weryfikacji', '', '| Pole | Pozycji |', '|---|---|']
    L += [f'| {k} | {v} |' for k, v in pola.most_common()]
    L += ['', 'Szczegóły: `dev/dane/do-weryfikacji.csv` albo zakładka „Do weryfikacji” w `import-do-arkusza.xlsx`.', '']
    with open(os.path.join(OUT, 'raport-konwersji.md'), 'w', encoding='utf-8') as f:
        f.write('\n'.join(L))

    print(f'Produkty: {len(produkty)}, zestawy: {len(zestawy)}, do weryfikacji: {len(weryf)}')
    print('Rodziny aktywnych:', dict(rodziny))
    print('Warianty:', {k: v for k, v in statystyka.items() if k.startswith('warianty')})


if __name__ == '__main__':
    main()
